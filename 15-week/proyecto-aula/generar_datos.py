"""
generar_datos.py
Datos sinteticos de DomiPan Huila S.A.S. (empresa ficticia)
Proyecto de aula · Inteligencia de Negocios 2026-B · Corhuila

Crea los dos archivos de origen que se cargan en Power BI:
  ventas_domipan.csv     -> tabla de hechos: una fila por producto dentro de un pedido
  maestros_domipan.xlsx  -> hojas Productos, Clientes, Sedes, Canales y Festivos

Periodo: 1 de octubre de 2025 al 30 de septiembre de 2026 (12 meses completos).
Requisitos: Python 3.10 o superior y openpyxl  (pip install openpyxl)
La SEMILLA fija hace que cada vez que se ejecute salgan exactamente los mismos datos.
"""

import csv
import random
from bisect import bisect_right
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

SEMILLA = 20251001
random.seed(SEMILLA)

INICIO = date(2025, 10, 1)
FIN = date(2026, 9, 30)
DIAS_PERIODO = (FIN - INICIO).days + 1


# =====================================================================
# 1. TABLAS MAESTRAS (dimensiones)
# =====================================================================

# id, nombre, zona, tipo, apertura, cierre
SEDES = [
    ("SED01", "DomiPan Centro", "Centro", "Punto de venta con planta de produccion", "06:00", "20:00"),
    ("SED02", "DomiPan Norte", "Norte", "Punto de venta", "06:00", "21:00"),
    ("SED03", "DomiPan Oriente", "Oriente", "Punto de venta", "06:00", "21:00"),
]

# id, nombre, tipo, comision (% de la venta), recargo de precio (% sobre el precio base)
CANALES = [
    ("CAN01", "Presencial", "Mostrador", 0.00, 0.00),
    ("CAN02", "Domicilio WhatsApp", "Domicilio", 0.05, 0.00),
    ("CAN03", "App movil", "Domicilio", 0.20, 0.10),
]

# Pedidos promedio por dia, por sede y canal (antes de aplicar temporadas)
PEDIDOS_BASE = {
    "SED01": {"CAN01": 45, "CAN02": 6, "CAN03": 4},
    "SED02": {"CAN01": 35, "CAN02": 10, "CAN03": 7},
    "SED03": {"CAN01": 25, "CAN02": 5, "CAN03": 3},
}

# id, nombre, categoria, subcategoria, unidad, precio base (COP), costo base (COP), tipico del Huila,
# peso de venta en cada franja: (desayuno, media manana, almuerzo, onces, noche)
PRODUCTOS = [
    ("PRD001", "Pandebono", "Panaderia", "Pan de queso y almidon", "Unidad", 2000, 820, "No", (10, 6, 3, 9, 4)),
    ("PRD002", "Almojabana", "Panaderia", "Pan de queso y almidon", "Unidad", 2000, 860, "No", (8, 5, 2, 7, 3)),
    ("PRD003", "Pan de queso", "Panaderia", "Pan de queso y almidon", "Unidad", 2500, 1050, "No", (5, 4, 2, 5, 3)),
    ("PRD004", "Pan frances", "Panaderia", "Pan de sal", "Unidad", 600, 230, "No", (12, 3, 2, 6, 10)),
    ("PRD005", "Pan aliñado", "Panaderia", "Pan de sal", "Unidad", 1500, 560, "No", (6, 2, 1, 4, 6)),
    ("PRD006", "Mogolla integral", "Panaderia", "Pan de sal", "Unidad", 1200, 470, "No", (5, 2, 1, 3, 3)),
    ("PRD007", "Pan tajado integral", "Panaderia", "Pan de sal", "Paquete", 8500, 4300, "No", (3, 1, 1, 2, 3)),
    ("PRD008", "Croissant de mantequilla", "Panaderia", "Hojaldre", "Unidad", 3500, 1650, "No", (4, 4, 1, 3, 1)),
    ("PRD009", "Roscon de arequipe", "Panaderia", "Pan dulce", "Unidad", 3000, 1200, "No", (3, 3, 1, 5, 2)),
    ("PRD010", "Mantecada", "Panaderia", "Pan dulce", "Unidad", 2500, 950, "No", (2, 3, 1, 4, 2)),
    ("PRD011", "Buñuelo", "Panaderia", "Pan de queso y almidon", "Unidad", 1800, 700, "No", (3, 2, 1, 3, 2)),
    ("PRD012", "Achiras del Huila x12", "Panaderia", "Tipicos del Huila", "Paquete", 7000, 3200, "Si", (1, 2, 1, 3, 2)),
    ("PRD013", "Pan de arroz", "Panaderia", "Tipicos del Huila", "Unidad", 1800, 720, "Si", (2, 2, 1, 3, 1)),
    ("PRD014", "Insulso huilense", "Panaderia", "Tipicos del Huila", "Unidad", 2500, 1000, "Si", (1, 1, 1, 2, 1)),
    ("PRD015", "Tinto", "Cafeteria", "Bebidas calientes", "Vaso", 1500, 290, "No", (14, 6, 4, 6, 2)),
    ("PRD016", "Cafe con leche", "Cafeteria", "Bebidas calientes", "Vaso", 3000, 880, "No", (9, 3, 1, 4, 2)),
    ("PRD017", "Capuchino", "Cafeteria", "Bebidas calientes", "Vaso", 5500, 1750, "No", (3, 3, 2, 3, 1)),
    ("PRD018", "Chocolate caliente", "Cafeteria", "Bebidas calientes", "Taza", 4000, 1350, "No", (5, 1, 1, 5, 3)),
    ("PRD019", "Aromatica", "Cafeteria", "Bebidas calientes", "Vaso", 2000, 380, "No", (2, 1, 2, 2, 3)),
    ("PRD020", "Jugo natural", "Cafeteria", "Bebidas frias", "Vaso", 5500, 2100, "No", (3, 4, 5, 3, 1)),
    ("PRD021", "Avena fria", "Cafeteria", "Bebidas frias", "Vaso", 4000, 1500, "No", (2, 2, 3, 3, 1)),
    ("PRD022", "Empanada de carne", "Cafeteria", "Comidas", "Unidad", 2800, 1150, "No", (1, 5, 6, 3, 1)),
    ("PRD023", "Sandwich de jamon y queso", "Cafeteria", "Comidas", "Unidad", 8500, 4100, "No", (1, 3, 5, 2, 1)),
    ("PRD024", "Desayuno completo", "Cafeteria", "Comidas", "Plato", 16000, 7600, "No", (4, 1, 0.2, 0, 0)),
    ("PRD025", "Porcion torta de chocolate", "Reposteria", "Tortas por porcion", "Porcion", 7000, 2900, "No", (0.5, 1, 1, 4, 3)),
    ("PRD026", "Porcion tres leches", "Reposteria", "Tortas por porcion", "Porcion", 7500, 3000, "No", (0.5, 1, 1, 4, 3)),
    ("PRD027", "Porcion cheesecake de maracuya", "Reposteria", "Tortas por porcion", "Porcion", 9500, 4300, "No", (0.3, 1, 1, 3, 3)),
    ("PRD028", "Milhoja", "Reposteria", "Postres individuales", "Unidad", 5500, 2200, "No", (0.5, 1, 1, 3, 2)),
    ("PRD029", "Postre de natas", "Reposteria", "Postres individuales", "Vaso", 6500, 2600, "No", (0.3, 1, 1, 2, 2)),
    ("PRD030", "Brownie", "Reposteria", "Postres individuales", "Unidad", 4500, 1700, "No", (0.5, 1, 1, 3, 2)),
    ("PRD031", "Torta de cumpleaños 1 libra", "Reposteria", "Tortas enteras", "Unidad", 48000, 21500, "No", (0.1, 0.3, 0.3, 0.6, 0.5)),
    ("PRD032", "Torta envinada 1 libra", "Reposteria", "Tortas enteras", "Unidad", 42000, 17500, "No", (0.05, 0.2, 0.2, 0.3, 0.3)),
]

# Nombres y categorias con tildes, como deben quedar despues del ETL (el script usa
# nombres sin tilde internamente para evitar errores de escritura).
NOMBRE_LIMPIO = {
    "Almojabana": "Almojábana", "Pan frances": "Pan francés", "Roscon de arequipe": "Roscón de arequipe",
    "Cafe con leche": "Café con leche", "Aromatica": "Aromática", "Avena fria": "Avena fría",
    "Sandwich de jamon y queso": "Sándwich de jamón y queso", "Porcion torta de chocolate": "Porción torta de chocolate",
    "Porcion tres leches": "Porción tres leches", "Porcion cheesecake de maracuya": "Porción cheesecake de maracuyá",
}
CATEGORIA_LIMPIA = {"Panaderia": "Panadería", "Cafeteria": "Cafetería", "Reposteria": "Repostería"}
SUBCATEGORIA_LIMPIA = {"Pan de queso y almidon": "Pan de queso y almidón", "Tipicos del Huila": "Típicos del Huila",
                       "Bebidas frias": "Bebidas frías", "Tortas por porcion": "Tortas por porción"}
UNIDAD_LIMPIA = {"Porcion": "Porción"}

# Festivos de Colombia en el periodo (Ley 51 de 1983, traslado a lunes)
FESTIVOS = [
    (date(2025, 10, 13), "Día de la Raza"), (date(2025, 11, 3), "Todos los Santos"),
    (date(2025, 11, 17), "Independencia de Cartagena"), (date(2025, 12, 8), "Inmaculada Concepción"),
    (date(2025, 12, 25), "Navidad"), (date(2026, 1, 1), "Año Nuevo"), (date(2026, 1, 12), "Reyes Magos"),
    (date(2026, 3, 23), "San José"), (date(2026, 4, 2), "Jueves Santo"), (date(2026, 4, 3), "Viernes Santo"),
    (date(2026, 5, 1), "Día del Trabajo"), (date(2026, 5, 18), "Ascensión del Señor"),
    (date(2026, 6, 8), "Corpus Christi"), (date(2026, 6, 15), "Sagrado Corazón"),
    (date(2026, 6, 29), "San Pedro y San Pablo"), (date(2026, 7, 20), "Independencia de Colombia"),
    (date(2026, 8, 7), "Batalla de Boyacá"), (date(2026, 8, 17), "Asunción de la Virgen"),
]
ES_FESTIVO = {f for f, _ in FESTIVOS}


# =====================================================================
# 2. CLIENTES (solo codigos y datos generales: nada personal)
# =====================================================================
# C0000 = consumidor final (el que compra en mostrador sin registrarse).
# C0001 a C0650 = clientes registrados para domicilios o tarjeta de puntos.
COMUNAS_POR_SEDE = {"SED01": [4, 6, 7], "SED02": [1, 2, 9], "SED03": [3, 5, 8, 10]}

clientes = []  # (id, tipo, comuna, fecha_registro, canal_registro, sede_cercana)
for i in range(1, 651):
    es_empresa = i % 22 == 0                      # unas 30 empresas (oficinas que piden desayunos)
    sede = random.choices(["SED01", "SED02", "SED03"], weights=[3, 4, 3])[0]
    if es_empresa:
        sede = random.choices(["SED01", "SED02", "SED03"], weights=[6, 2, 2])[0]
    comuna = random.choice(COMUNAS_POR_SEDE[sede])
    if i <= 400:   # clientes antiguos: se registraron antes del periodo
        registro = date(2023, 1, 1) + timedelta(days=random.randint(0, (INICIO - date(2023, 1, 1)).days - 1))
        canal_reg = random.choices(["Tienda", "WhatsApp", "App"], weights=[50, 40, 10])[0]
    else:          # clientes nuevos: llegan durante el periodo, sobre todo por la app
        registro = INICIO + timedelta(days=int(random.triangular(0, DIAS_PERIODO - 1, DIAS_PERIODO - 1)))
        canal_reg = random.choices(["Tienda", "WhatsApp", "App"], weights=[10, 35, 55])[0]
    if es_empresa:
        canal_reg = random.choice(["Tienda", "WhatsApp"])
    clientes.append((f"C{i:04d}", "Empresa" if es_empresa else "Persona", comuna, registro, canal_reg, sede))

# Listas ordenadas por fecha de registro, para elegir solo clientes ya registrados el dia del pedido
def _ordenar(lista):
    lista = sorted(lista, key=lambda c: c[3])
    return lista, [c[3] for c in lista]

PERSONAS_POR_SEDE = {s: _ordenar([c for c in clientes if c[1] == "Persona" and c[5] == s]) for s in COMUNAS_POR_SEDE}
PERSONAS_TODAS = _ordenar([c for c in clientes if c[1] == "Persona"])
EMPRESAS_POR_SEDE = {s: _ordenar([c for c in clientes if c[1] == "Empresa" and c[5] == s]) for s in COMUNAS_POR_SEDE}


def cliente_registrado(lista_y_fechas, dia):
    """Elige al azar un cliente que ya estaba registrado ese dia."""
    lista, fechas = lista_y_fechas
    k = bisect_right(fechas, dia)
    return random.choice(lista[:k])[0] if k > 0 else None


# =====================================================================
# 3. REGLAS DE NEGOCIO (para que los datos se parezcan a una panaderia real)
# =====================================================================
FRANJAS = ["Desayuno", "Media mañana", "Almuerzo", "Onces", "Noche"]

def franja_de(hora):
    if hora < 10: return 0
    if hora < 12: return 1
    if hora < 15: return 2
    if hora < 18: return 3
    return 4

# Peso de cada hora del dia segun el canal (de 6 a. m. a 8 p. m.)
PESO_HORA = {
    "CAN01": {6: 9, 7: 12, 8: 10, 9: 6, 10: 5, 11: 4, 12: 4, 13: 4, 14: 4, 15: 6, 16: 8, 17: 9, 18: 7, 19: 5, 20: 3},
    "CAN02": {6: 3, 7: 8, 8: 9, 9: 6, 10: 4, 11: 4, 12: 3, 13: 3, 14: 3, 15: 5, 16: 7, 17: 7, 18: 6, 19: 4, 20: 2},
    "CAN03": {6: 0, 7: 3, 8: 5, 9: 4, 10: 3, 11: 3, 12: 3, 13: 3, 14: 3, 15: 6, 16: 8, 17: 9, 18: 9, 19: 7, 20: 4},
}

# Efecto del dia de la semana (lunes=0 ... domingo=6). Un festivo se comporta como domingo.
DIA_SEMANA_MOSTRADOR = [0.90, 0.90, 0.95, 0.95, 1.05, 1.15, 1.00]
DIA_SEMANA_DOMICILIO = [0.90, 0.90, 0.90, 0.95, 1.05, 1.30, 1.40]

# Efecto del mes: diciembre alto (Navidad), enero bajo (cuesta de enero), junio alto (San Pedro)
FACTOR_MES = {10: 1.00, 11: 1.00, 12: 1.35, 1: 0.85, 2: 0.95, 3: 1.00, 4: 0.97,
              5: 1.05, 6: 1.12, 7: 1.05, 8: 1.00, 9: 1.03}

# Que tanto se pide cada subcategoria por cada canal (1 = igual que en mostrador)
PREFERENCIA_CANAL = {
    "CAN02": {"Pan de sal": 1.5, "Pan de queso y almidon": 1.3, "Bebidas calientes": 0.5,
              "Comidas": 1.3, "Tortas enteras": 4.0},
    "CAN03": {"Bebidas calientes": 0.3, "Bebidas frias": 0.8, "Tortas por porcion": 2.5,
              "Postres individuales": 2.5, "Tortas enteras": 4.0, "Comidas": 1.2,
              "Pan de sal": 0.8, "Pan de queso y almidon": 0.8, "Hojaldre": 0.8, "Pan dulce": 0.8},
}

PANES_PEQUENOS = {"PRD001", "PRD002", "PRD003", "PRD005", "PRD006", "PRD011", "PRD013"}


def factor_del_dia(dia, sede, canal):
    """Multiplica los pedidos base segun dia, mes, quincena, tendencia y eventos especiales."""
    dsem = 6 if dia in ES_FESTIVO else dia.weekday()
    f = (DIA_SEMANA_MOSTRADOR if canal == "CAN01" else DIA_SEMANA_DOMICILIO)[dsem]
    if sede == "SED01" and canal == "CAN01" and dsem == 6:
        f *= 0.55                                   # el Centro vive de oficinas: domingos y festivos flojos
    f *= FACTOR_MES[dia.month]
    if dia.day in (1, 15, 16, 30, 31):
        f *= 1.08                                   # dias de pago de quincena
    if canal == "CAN03":
        f *= 0.75 + 0.60 * (dia - INICIO).days / (DIAS_PERIODO - 1)   # la app crece todo el año
    if date(2025, 12, 16) <= dia <= date(2025, 12, 23):
        f *= 1.15                                   # novenas de aguinaldo
    if dia in (date(2025, 12, 24), date(2025, 12, 31)):
        f *= 1.25
    if dia in (date(2025, 12, 25), date(2026, 1, 1)):
        f *= 0.40                                   # se abre medio dia
    if sede == "SED01" and date(2026, 1, 2) <= dia <= date(2026, 1, 12):
        f *= 0.75                                   # oficinas del centro en vacaciones
    if date(2026, 6, 19) <= dia <= date(2026, 7, 1):
        f *= 1.18                                   # Festival del Bambuco y San Pedro en Neiva
    if date(2026, 5, 8) <= dia <= date(2026, 5, 10):
        f *= 1.10                                   # Dia de la Madre
    if date(2026, 9, 17) <= dia <= date(2026, 9, 19):
        f *= 1.08                                   # Amor y Amistad
    return f


def ajuste_producto(prod, dia):
    """Productos que suben en fechas especiales y fines de semana."""
    pid, cat, sub = prod[0], prod[2], prod[3]
    f = 1.0
    if dia.month == 12 and pid == "PRD011": f *= 4.0               # buñuelos en Navidad
    if dia.month == 12 and pid == "PRD032": f *= 5.0               # torta envinada en Navidad
    if dia in (date(2025, 12, 24), date(2025, 12, 31)) and sub == "Tortas enteras": f *= 3.0
    if date(2026, 6, 19) <= dia <= date(2026, 7, 1) and sub == "Tipicos del Huila": f *= 2.5
    if date(2026, 5, 8) <= dia <= date(2026, 5, 10):
        if sub == "Tortas enteras": f *= 5.0
        elif cat == "Reposteria": f *= 1.8
    if date(2026, 9, 17) <= dia <= date(2026, 9, 19) and cat == "Reposteria": f *= 2.0
    if dia.weekday() >= 5 and cat == "Reposteria": f *= 1.4
    if dia.weekday() == 6 and pid == "PRD024": f *= 1.5             # desayuno completo de domingo
    return f


def redondear_50(valor):
    return int(round(valor / 50.0)) * 50


def precio_venta(prod, dia, canal):
    """Precio base; sube 6 % desde el 1 de febrero de 2026; la app cobra un recargo del 10 %."""
    precio = prod[5]
    if dia >= date(2026, 2, 1):
        precio = redondear_50(precio * 1.06)
    if canal == "CAN03":
        precio = redondear_50(precio * 1.10)
    return precio


def costo_unitario(prod, dia):
    """Costo base; los insumos suben 8 % desde el 1 de enero de 2026 (un mes antes que los precios)."""
    costo = prod[6]
    if dia >= date(2026, 1, 1):
        costo = int(round(costo * 1.08, -1))
    return costo


def cantidad(prod, canal, es_empresa):
    pid, sub = prod[0], prod[3]
    if sub == "Tortas enteras" or prod[4] == "Paquete":
        q = random.choices([1, 2], weights=[92, 8])[0]
    elif pid == "PRD004":
        q = random.choices([1, 2, 3, 4, 5, 6, 8, 10], weights=[8, 10, 10, 14, 18, 14, 14, 12])[0]
    elif pid in PANES_PEQUENOS:
        q = random.choices([1, 2, 3, 4, 5, 6], weights=[30, 30, 18, 10, 7, 5])[0]
    else:
        q = random.choices([1, 2, 3], weights=[70, 22, 8])[0]
    if canal != "CAN01" and (pid == "PRD004" or pid in PANES_PEQUENOS):
        q = int(q * random.choice([1, 1, 1.5, 2]))   # a domicilio se pide pan para toda la casa
    if es_empresa and sub not in ("Tortas enteras",):
        q = min(q * random.randint(4, 12), 60)         # pedidos de oficina, maximo 60 por producto
    return max(q, 1)


# =====================================================================
# 4. GENERACION DE PEDIDOS
# =====================================================================
filas = []          # filas limpias: una por producto dentro de un pedido
num_pedido = 0
dia = INICIO
while dia <= FIN:
    pedidos_del_dia = []
    for sede, apertura, cierre in [(s[0], s[4], s[5]) for s in SEDES]:
        ultima_hora = int(cierre[:2]) - 1
        for canal in ("CAN01", "CAN02", "CAN03"):
            esperado = PEDIDOS_BASE[sede][canal] * factor_del_dia(dia, sede, canal)
            n = max(0, round(esperado * random.uniform(0.85, 1.15)))
            horas = [h for h in PESO_HORA[canal] if h <= ultima_hora]
            pesos = [PESO_HORA[canal][h] for h in horas]
            if dia in (date(2025, 12, 25), date(2026, 1, 1)):     # medio dia: de 7 a. m. a 1 p. m.
                horas, pesos = [h for h in range(7, 13)], [5, 6, 5, 4, 3, 2]
            for _ in range(n):
                hora = random.choices(horas, weights=pesos)[0]
                minuto = random.randint(0, 59)
                pedidos_del_dia.append((hora, minuto, sede, canal))
    pedidos_del_dia.sort()

    for hora, minuto, sede, canal in pedidos_del_dia:
        num_pedido += 1
        id_pedido = f"P{num_pedido:06d}"
        franja = franja_de(hora)

        # Cliente del pedido
        es_empresa = False
        if canal == "CAN01":
            if random.random() < 0.15:
                cli = cliente_registrado(PERSONAS_POR_SEDE[sede], dia)   # cliente con tarjeta de puntos
            else:
                cli = "C0000"                                            # consumidor final
        else:
            if canal == "CAN02" and dia.weekday() < 5 and franja <= 1 and random.random() < 0.30:
                cli = cliente_registrado(EMPRESAS_POR_SEDE[sede], dia)
                es_empresa = cli is not None
            else:
                cli = None
            if cli is None:
                fuente = PERSONAS_POR_SEDE[sede] if random.random() < 0.8 else PERSONAS_TODAS
                cli = cliente_registrado(fuente, dia)
        cli = cli or "C0000"

        # Cuantos productos distintos lleva el pedido
        if canal == "CAN01":
            n_items = random.choices([1, 2, 3, 4], weights=[40, 35, 18, 7])[0]
        else:
            n_items = random.choices([1, 2, 3, 4, 5], weights=[15, 30, 30, 17, 8])[0]

        # Que productos: segun la franja, el canal, la fecha y el tipo de cliente
        candidatos = list(PRODUCTOS)
        pesos = []
        for p in candidatos:
            w = p[8][franja] * ajuste_producto(p, dia)
            w *= PREFERENCIA_CANAL.get(canal, {}).get(p[3], 1.0)
            if es_empresa and p[0] in ("PRD024", "PRD015", "PRD016", "PRD001", "PRD002", "PRD020"):
                w *= 6
            pesos.append(w)
        elegidos = []
        for _ in range(n_items):
            if sum(pesos) <= 0:
                break
            idx = random.choices(range(len(candidatos)), weights=pesos)[0]
            elegidos.append(candidatos.pop(idx))
            pesos.pop(idx)

        for linea, prod in enumerate(elegidos, start=1):
            filas.append({
                "id_pedido": id_pedido, "linea": linea, "fecha": dia, "hora": f"{hora:02d}:{minuto:02d}",
                "id_sede": sede, "id_canal": canal, "id_cliente": cli, "id_producto": prod[0],
                "cantidad": cantidad(prod, canal, es_empresa),
                "precio_unitario": precio_venta(prod, dia, canal),
                "costo_unitario": costo_unitario(prod, dia),
            })
    dia += timedelta(days=1)


# =====================================================================
# 5. PROBLEMAS DE CALIDAD A PROPOSITO (para que el ETL tenga sentido)
# =====================================================================
# Problema 1: pedidos duplicados. Al sincronizar, unos domicilios quedaron registrados dos veces.
pedidos_domicilio = sorted({f["id_pedido"] for f in filas if f["id_canal"] != "CAN01"})
duplicados = set(random.sample(pedidos_domicilio, k=round(len(pedidos_domicilio) * 0.015)))

# Problemas 2 y 3: la app exporta con otro formato -> fecha dd/mm/aaaa y precio como texto "$2.200".
def fila_para_csv(f):
    if f["id_canal"] == "CAN03":
        fecha = f["fecha"].strftime("%d/%m/%Y")
        precio = "$" + f"{f['precio_unitario']:,}".replace(",", ".")
    else:
        fecha = f["fecha"].isoformat()
        precio = str(f["precio_unitario"])
    return [f["id_pedido"], f["linea"], fecha, f["hora"], f["id_sede"], f["id_canal"], f["id_cliente"],
            f["id_producto"], f["cantidad"], precio, f["costo_unitario"]]

ENCABEZADO = ["id_pedido", "linea", "fecha", "hora", "id_sede", "id_canal", "id_cliente",
              "id_producto", "cantidad", "precio_unitario", "costo_unitario"]

filas_csv = []
for i, f in enumerate(filas):
    filas_csv.append(fila_para_csv(f))
    es_ultima_linea = i + 1 == len(filas) or filas[i + 1]["id_pedido"] != f["id_pedido"]
    if es_ultima_linea and f["id_pedido"] in duplicados:
        # al terminar un pedido duplicado se repiten todas sus lineas (la ultima linea = cuantas tiene)
        lineas_del_pedido = filas[i + 1 - f["linea"]: i + 1]
        filas_csv.extend(fila_para_csv(x) for x in lineas_del_pedido)

with open("ventas_domipan.csv", "w", newline="", encoding="utf-8") as archivo:
    escritor = csv.writer(archivo)
    escritor.writerow(ENCABEZADO)
    escritor.writerows(filas_csv)


# Problema 4: en la hoja Productos, nombres y categorias se escribieron de formas distintas.
CATEGORIA_SUCIA = {"PRD003": "Panaderia", "PRD008": "PANADERÍA", "PRD012": "panadería ",
                   "PRD017": "Cafeteria", "PRD027": "Reposteria", "PRD030": "REPOSTERÍA",
                   "PRD031": "repostería "}
NOMBRE_SUCIO = {"PRD002": "  almojabana", "PRD003": "PAN DE QUESO", "PRD008": "croissant de mantequilla ",
                "PRD016": "Cafe Con Leche", "PRD026": "porcion tres leches  "}


# =====================================================================
# 6. ARCHIVO DE EXCEL CON LAS TABLAS MAESTRAS
# =====================================================================
libro = Workbook()
negrita = Font(name="Arial", bold=True)
normal = Font(name="Arial")


def hoja(nombre, encabezados, datos, formatos=None):
    ws = libro.create_sheet(nombre)
    ws.append(encabezados)
    for fila in datos:
        ws.append(list(fila))
    for fila in ws.iter_rows():
        for celda in fila:
            celda.font = negrita if celda.row == 1 else normal
    for col, fmt in (formatos or {}).items():
        for celda in ws[col][1:]:
            celda.number_format = fmt
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or "")) for c in col) + 3
    ws.freeze_panes = "A2"


libro.remove(libro.active)

hoja("Productos",
     ["id_producto", "producto", "categoria", "subcategoria", "unidad_venta", "precio_base", "costo_base", "es_tipico_huila"],
     [(p[0],
       NOMBRE_SUCIO.get(p[0], NOMBRE_LIMPIO.get(p[1], p[1])),
       CATEGORIA_SUCIA.get(p[0], CATEGORIA_LIMPIA[p[2]]),
       SUBCATEGORIA_LIMPIA.get(p[3], p[3]),
       UNIDAD_LIMPIA.get(p[4], p[4]),
       p[5], p[6],
       "Sí" if p[7] == "Si" else "No") for p in PRODUCTOS],
     {"F": "#,##0", "G": "#,##0"})

hoja("Clientes",
     ["id_cliente", "tipo_cliente", "comuna", "fecha_registro", "canal_registro"],
     [("C0000", "Consumidor final", "No aplica", None, "No aplica")] +
     [(c[0], c[1], f"Comuna {c[2]}", c[3], c[4]) for c in clientes],
     {"D": "yyyy-mm-dd"})

hoja("Sedes",
     ["id_sede", "sede", "zona", "ciudad", "tipo_sede", "hora_apertura", "hora_cierre"],
     [(s[0], s[1], s[2], "Neiva", s[3].replace("produccion", "producción"), s[4], s[5]) for s in SEDES])

hoja("Canales",
     ["id_canal", "canal", "tipo_canal", "comision_pct", "recargo_precio_pct"],
     [(c[0], c[1].replace("movil", "móvil"), c[2], c[3], c[4]) for c in CANALES],
     {"D": "0%", "E": "0%"})

hoja("Festivos", ["fecha", "festivo"], FESTIVOS, {"A": "yyyy-mm-dd"})

libro.save("maestros_domipan.xlsx")

print(f"Pedidos generados: {num_pedido:,}")
print(f"Filas limpias: {len(filas):,}  |  Filas en el CSV (con duplicados): {len(filas_csv):,}")
